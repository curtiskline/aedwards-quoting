"""T531 cell-for-cell verification of the ported OmegaWrap engine.

Two independent checks against docs/reference/maxim-handoff/
AE_Composite_Calculator_v2.06_072926_NO_TimeWarp.xlsm:

1. ANCHOR: the engine vs the workbook's own cached values (openpyxl
   data_only=True) at the exact inputs the file was saved with. This is
   Excel's real output, bit for bit.

2. GRID: the workbook's formula text is copied VERBATIM (openpyxl, no
   retyping) into a minimal workbook, evaluated independently by the
   `formulas` library over a deterministic grid of inputs, and compared
   with the engine cell by cell. The copied cells are the full J-column
   chains of 'Corrosion Calculations' and 'Mechanical Calculations' plus
   their gating flags; 'OW Variables' is carried over as cached values
   (it is a constant table).

Cells NOT copied (recorded in the report):
  - J62 / J60 Spiral wrap length: VBA UDF RepairWrapLengthSpiral, out of
    scope (application technique is pinned to Layer over Layer).
  - I19 install-temperature flag: contains #REF! errors in the shipped
    workbook and never feeds a number.
  - J35 / J33 'Service Factor Test' text and the M/N error-message blocks:
    display text only.

Run:  venv/bin/python3 scripts/verify_omegawrap_xlsm.py
Writes docs/research/omegawrap-formula-port-T531-verification.md and exits
non-zero on any divergence.
"""

from __future__ import annotations

import os
import sys
import tempfile
import warnings
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

warnings.filterwarnings("ignore")

import formulas  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

from src.app.omegawrap_formulas import (  # noqa: E402
    RepairOutputs,
    corrosion_outputs,
    mechanical_outputs,
)

XLSM = os.path.join(
    os.path.dirname(__file__), "..", "docs", "reference", "maxim-handoff",
    "AE_Composite_Calculator_v2.06_072926_NO_TimeWarp.xlsm",
)
REPORT = os.path.join(
    os.path.dirname(__file__), "..", "docs", "research",
    "omegawrap-formula-port-T531-verification.md",
)
REL_TOL = 1e-9

# Verbatim-copied formula cells per calculations sheet.
CORR_CELLS = (
    ["J%d" % r for r in (10, 11, 12, 13, 14, 15, 16, 17, 20, 21, 22, 23, 24,
                         25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 36, 38, 39,
                         40, 41, 42, 43, 44, 45, 46, 47, 50, 51, 52, 53, 55,
                         56, 57, 59, 60, 65, 66, 67, 68, 69, 70, 71, 72, 73,
                         74)]
    + ["I20", "I36", "I55"]
)
MECH_CELLS = (
    ["J%d" % r for r in (10, 11, 12, 13, 14, 15, 16, 17, 20, 21, 22, 23, 24,
                         25, 26, 27, 28, 29, 30, 31, 32, 34, 36, 37, 38, 39,
                         40, 41, 42, 43, 44, 45, 46, 49, 50, 51, 53, 54, 55,
                         57, 58, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72)]
    + ["I34", "I53"]
)

# engine field -> workbook output cell, per sheet. The gated J38-J47 /
# J36-J45 cells are used (not the raw calc rows) so the depth gate is
# verified too.
CORR_MAP = {
    "maop_psig": "J16", "derated_pressure_psig": "J17",
    "number_of_layers": "J38", "repair_thickness_in": "J39",
    "finished_od_in": "J40", "fabric_length_per_set_ft": "J41",
    "linear_length_of_repair_in": "J42", "total_fabric_length_ft": "J43",
    "resin_per_set_qt_per_ft": "J44", "total_resin_qt": "J45",
    "primer_resin_qt": "J46", "putty_pt": "J47", "wrap_segments": "J56",
    "isolation_wrap_per_ft": "J71", "isolation_wrap_total": "J72",
    "isolation_wrap_resin_qt": "J73",
}
MECH_MAP = {
    "maop_psig": "J16", "derated_pressure_psig": "J17",
    "number_of_layers": "J36", "repair_thickness_in": "J37",
    "finished_od_in": "J38", "fabric_length_per_set_ft": "J39",
    "linear_length_of_repair_in": "J40", "total_fabric_length_ft": "J41",
    "resin_per_set_qt_per_ft": "J42", "total_resin_qt": "J43",
    "primer_resin_qt": "J44", "putty_pt": "J45", "wrap_segments": "J54",
    "isolation_wrap_per_ft": "J69", "isolation_wrap_total": "J70",
    "isolation_wrap_resin_qt": "J71",
}

GRADES = ["24,000", "A25", "A", "B", "X42", "X46", "X52", "X56", "X60",
          "X65", "X70"]
PCT_SMYS = [72, 60, 50, 40]  # class 1..4


def grid_points():
    """Deterministic input grid: a round-robin sweep over every axis, plus
    hand-picked edge cases. ~190 points per sheet."""
    ods = [4.5, 6.625, 8.625, 12.75, 16, 20, 24, 30, 36]
    wts = [0.154, 0.188, 0.25, 0.375, 0.5, 0.625]
    depths = [0.02, 0.05, 0.1, 0.15, 0.21]
    lengths = [1, 6, 11.5, 12, 24, 52.81, 120, 240]
    widths = [0.5, 2, 3.14, 6, 12]
    composites = ["E-Glass", "Carbon"]
    points = []
    for i in range(180):
        points.append(dict(
            od=ods[i % len(ods)], wt=wts[(i // 2) % len(wts)],
            grade=GRADES[(i // 3) % len(GRADES)],
            pct=PCT_SMYS[(i // 5) % len(PCT_SMYS)],
            depth=depths[(i // 7) % len(depths)],
            length=lengths[(i // 4) % len(lengths)],
            width=widths[(i // 6) % len(widths)],
            composite=composites[i % 2],
        ))
    # Edges: exact 12-multiple repair length (OD*WT/2 = 4 -> L_over = 5,
    # 2*5 + 2 = 12), depth == wall (corrosion gate), depth > wall, huge and
    # tiny pipes, layer count exactly at the minimum-3 boundary.
    points += [
        dict(od=8, wt=1, grade="X52", pct=72, depth=0.3, length=2, width=4, composite="Carbon"),
        dict(od=8, wt=1, grade="X52", pct=72, depth=0.3, length=2.0000001, width=4, composite="E-Glass"),
        dict(od=20, wt=0.25, grade="X70", pct=40, depth=0.25, length=52.81, width=6, composite="Carbon"),
        dict(od=20, wt=0.25, grade="X70", pct=40, depth=0.26, length=52.81, width=6, composite="E-Glass"),
        dict(od=4.5, wt=0.154, grade="24,000", pct=40, depth=0.01, length=1, width=0.5, composite="E-Glass"),
        dict(od=36, wt=0.625, grade="X70", pct=72, depth=0.6, length=240, width=12, composite="Carbon"),
        dict(od=12.75, wt=0.044, grade="A25", pct=50, depth=0.01, length=6, width=2, composite="Carbon"),
    ]
    return points


def build_mini_workbook() -> str:
    """Copy the formula chains verbatim from the xlsm into a fresh xlsx the
    `formulas` library can evaluate. Returns the temp file path."""
    src = openpyxl.load_workbook(XLSM)
    src_values = openpyxl.load_workbook(XLSM, data_only=True)
    mini = openpyxl.Workbook()
    mini.remove(mini.active)

    ow = mini.create_sheet("OW Variables")
    for row in src_values["OW Variables"].iter_rows(min_row=1, max_row=30,
                                                    max_col=35):
        for cell in row:
            if cell.value is not None:
                ow[cell.coordinate] = cell.value

    for sheet, input_cells, s_cells, calc_sheet, calc_cells in [
        ("Corrosion",
         {"I8": "Imperial", "E6": 70, "E21": 0.5, "E33": "Layer over Layer"},
         ["S4", "S5", "S6", "S13", "S15", "S16"],
         "Corrosion Calculations", CORR_CELLS),
        ("Mechanical",
         {"G5": "Imperial", "E6": 80, "E20": 0.5, "E32": "Layer over Layer"},
         ["S4", "S5", "S6", "S13", "S15", "S16"],
         "Mechanical Calculations", MECH_CELLS),
    ]:
        ws = mini.create_sheet(sheet)
        for coord, value in input_cells.items():
            ws[coord] = value
        for coord in s_cells:
            ws[coord] = src[sheet][coord].value  # verbatim formula
        cs = mini.create_sheet(calc_sheet)
        for coord in calc_cells:
            value = src[calc_sheet][coord].value
            if isinstance(value, ArrayFormula):
                value = value.text
            assert isinstance(value, str) and value.startswith("="), (
                f"{calc_sheet}!{coord} is not a formula: {value!r}")
            cs[coord] = value

    path = os.path.join(tempfile.mkdtemp(), "mini.xlsx")
    mini.save(path)
    return path


class MiniModel:
    def __init__(self, path: str):
        self.model = formulas.ExcelModel().loads(path).finish()
        solution = self.model.calculate()
        # Map (SHEET upper, CELL) -> the solver's exact key string.
        self.keys = {}
        for key in solution:
            if "!" in key and "]" in key:
                sheet = key.split("]")[1].split("'!")[0].strip("'").upper()
                cell = key.split("!")[-1]
                self.keys[(sheet, cell)] = key

    def key(self, sheet: str, cell: str) -> str:
        return self.keys[(sheet.upper(), cell)]

    def calculate(self, corr: dict, mech: dict) -> dict:
        inputs = {}
        for sheet, values in [("CORROSION", corr), ("MECHANICAL", mech)]:
            for cell, value in values.items():
                inputs[self.key(sheet, cell)] = value
        return self.model.calculate(inputs=inputs)

    @staticmethod
    def scalar(solution, key):
        value = solution[key].value
        try:
            return value[0, 0]
        except Exception:
            return value


def close(a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return str(a) == str(b)
    if a is None or b is None:
        return a is b
    return abs(float(a) - float(b)) <= REL_TOL * max(1.0, abs(float(a)), abs(float(b)))


def input_cells(point, sheet):
    grade = 24000 if point["grade"] == "24,000" else point["grade"]
    cells = {"E4": point["od"], "E5": point["wt"], "E7": grade,
             "E10": point["pct"], "E13": point["depth"],
             "E15": point["length"], "E16": point["width"]}
    cells["E19" if sheet == "Corrosion" else "E18"] = point["composite"]
    return cells


def run_engine(point, sheet):
    fn = corrosion_outputs if sheet == "Corrosion" else mechanical_outputs
    return fn(point["od"], point["wt"], point["grade"], point["pct"],
              point["depth"], point["length"], point["width"],
              composite=point["composite"])


def main() -> int:
    divergences, notes = [], []
    engine_fields = list(CORR_MAP)

    # ---- 1. ANCHOR: engine vs Excel's own cached values ----
    cached = openpyxl.load_workbook(XLSM, data_only=True)
    anchor_checks = 0
    anchor_specs = [
        ("Corrosion", "Corrosion Calculations", CORR_MAP,
         dict(od=20, wt=0.25, grade="X70", pct=40, depth=0.21, length=52.81,
              width=6, composite="Carbon")),
        ("Mechanical", "Mechanical Calculations", MECH_MAP,
         dict(od=30, wt=0.375, grade="X52", pct=72, depth=0.19, length=12,
              width=5, composite="Carbon")),
    ]
    for sheet, calc_sheet, cell_map, point in anchor_specs:
        out = run_engine(point, sheet)
        for field, cell in cell_map.items():
            expected = cached[calc_sheet][cell].value
            got = getattr(out, field)
            anchor_checks += 1
            if not close(got, expected):
                divergences.append(
                    f"ANCHOR {calc_sheet}!{cell} ({field}): Excel cached "
                    f"{expected!r}, engine {got!r}, inputs {point}")

    # ---- 2. GRID: engine vs `formulas` evaluating the verbatim chains ----
    print("Building mini workbook from verbatim formulas...")
    mini_path = build_mini_workbook()
    model = MiniModel(mini_path)
    points = grid_points()
    grid_checks = 0
    print(f"Evaluating {len(points)} grid points x 2 sheets...")
    for index, point in enumerate(points):
        solution = model.calculate(input_cells(point, "Corrosion"),
                                   input_cells(point, "Mechanical"))
        for sheet, calc_sheet, cell_map in [
            ("Corrosion", "Corrosion Calculations", CORR_MAP),
            ("Mechanical", "Mechanical Calculations", MECH_MAP),
        ]:
            try:
                out = run_engine(point, sheet)
            except ValueError:
                out = None  # corrosion depth gate: expect '-' in Excel
            for field, cell in cell_map.items():
                expected = model.scalar(solution, model.key(calc_sheet, cell))
                got = None if out is None else getattr(out, field)
                grid_checks += 1
                gated = cell in ("J38", "J39", "J41", "J42", "J43", "J44",
                                 "J45", "J47") and sheet == "Corrosion"
                if out is None:
                    if cell == "J46":
                        # Known workbook quirk: J46 (primer) gates on I21,
                        # an EMPTY cell, instead of I20 like every other
                        # output, so Excel still shows primer resin on an
                        # invalid repair. The engine intentionally gates it
                        # with the rest; here we verify Excel's exact
                        # behavior: J46 stays numeric under the depth gate.
                        import math as _math
                        llr = _math.ceil(
                            (2 * 2.5 * _math.sqrt(point["od"] * point["wt"] / 2)
                             + point["length"]) / 12) * 12
                        primer = ((_math.pi * (point["od"] + 0.06))
                                  - (_math.pi * point["od"])) * (llr + 4) / 57.75
                        if not close(primer, expected):
                            divergences.append(
                                f"GRID #{index} {calc_sheet}!J46 under depth "
                                f"gate: Excel {expected!r}, expected quirk "
                                f"value {primer!r}; inputs {point}")
                        continue
                    if gated and str(expected) != "-":
                        divergences.append(
                            f"GRID #{index} {calc_sheet}!{cell}: engine gated "
                            f"(depth>=wall) but Excel shows {expected!r}; "
                            f"inputs {point}")
                    continue
                if not close(got, expected):
                    divergences.append(
                        f"GRID #{index} {calc_sheet}!{cell} ({field}): "
                        f"formulas-eval {expected!r}, engine {got!r}, "
                        f"inputs {point}")
        if (index + 1) % 25 == 0:
            print(f"  {index + 1}/{len(points)} points, "
                  f"{len(divergences)} divergences")

    # ---- report ----
    os.makedirs(os.path.dirname(REPORT), exist_ok=True)
    status = "PASS — no divergences" if not divergences else (
        f"FAIL — {len(divergences)} divergence(s)")
    with open(REPORT, "w") as f:
        f.write(f"""# OmegaWrap formula port verification (T531)

Generated by scripts/verify_omegawrap_xlsm.py on {date.today()}.

**Result: {status}.**
{anchor_checks} anchor checks (engine vs Excel's cached values at the
workbook's shipped inputs) and {grid_checks} grid checks (engine vs the
`formulas` library independently evaluating the workbook's verbatim formula
text over {len(points)} input points x 2 sheets x both composite types).

## What was ported
- 'Corrosion Calculations' full J-column chain (MAOP, PCC-2 Eq 12/17/19,
  layers with minimum 3, wrap segments, repair thickness, fabric lengths,
  resin, primer, putty, isolation wrap) with the depth>=wall gate (I20).
- 'Mechanical Calculations' full J-column chain (layers = 1.5 x wall /
  layer thickness with minimum 3, same tail; resin values ROUNDUP'd to 2
  decimals as that sheet does). The workbook has no depth gate on this
  sheet (its I20 is empty) and none was added.
- 'OW Variables' constant tables: pipe grades (SMYS/UTS), composite
  properties, unit conversions, putty density, error messages.
- Wheel Count Calculator G2->H2 (+12 in isolation reach) was already in
  the app (T525).

## Pinned defaults (workbook inputs the simple page does not ask for)
- Service Factor 0.5; Application Technique 'Layer over Layer' (overlap
  0%); Imperial units; install temperature 70/80 F (drives warning flags
  only, never a number in these chains).

## Out of scope / not ported
- Spiral application technique (VBA UDF RepairWrapLengthSpiral,
  Corrosion J62 / Mechanical J60): engine is pinned to Layer over Layer.
- Install-temperature warning flags (I19): formula contains #REF! errors
  in the shipped workbook; display-only.
- AE part-number BOM lookups (Corrosion/Mechanical rows 34-40 array
  formulas into Materials/BOM sheets), girth-weld sheets, Wrinkle, Cyclic
  Loading, pricing sheets: separate features, not this task's outputs.

## Documented intentional divergence (the only one)
- 'Corrosion Calculations'!J46 (Primer Resin Required) gates on I21,
  which is an EMPTY cell, instead of I20 like every other output row —
  almost certainly an off-by-one in the workbook's gate reference. Result
  in Excel: when defect depth >= wall thickness, every output shows '-'
  EXCEPT primer resin, which still shows a number. The ported engine
  gates ALL outputs together (an invalid repair shows the error message
  and no numbers). The harness still verifies Excel's exact J46 quirk
  value at every gated grid point, so this difference is checked, not
  ignored. Question that would have gone to Maxim (answered by
  inference per PM direction): was J46's I21 gate intentional? Inferred
  no — a lone primer quantity on a repair the sheet itself calls invalid
  ('Defect depth must be less than wall thickness') has no use.

## Divergences
""")
        if divergences:
            for d in divergences:
                f.write(f"- {d}\n")
        else:
            f.write("None.\n")
    print(f"\n{status}")
    print(f"Report: {REPORT}")
    return 1 if divergences else 0


if __name__ == "__main__":
    sys.exit(main())
